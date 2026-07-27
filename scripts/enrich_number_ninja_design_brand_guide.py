from __future__ import annotations

from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
BRAND = ROOT / "branding"
DOCX = BRAND / "brandbook" / "NumberNinjaDesign Brand Guide.docx"
BG = "#07090C"
SURFACE = "#11151B"
SURFACE_2 = "#171D25"
ACCENT = "#00E891"
ACCENT_2 = "#6EE7FF"
TEXT = "#F3F5F7"
MUTED = "#8B96A5"
LINE = "#26303C"


def rgb(value: str) -> tuple[int, int, int]:
    value = value.lstrip("#")
    return tuple(int(value[i:i + 2], 16) for i in (0, 2, 4))


def ft(size: int, bold: bool = False, display: bool = False):
    if display:
        path = Path(r"C:\Windows\Fonts\bahnschrift.ttf")
    elif bold:
        path = Path(r"C:\Windows\Fonts\arialbd.ttf")
    else:
        path = Path(r"C:\Windows\Fonts\arial.ttf")
    return ImageFont.truetype(str(path), size)


def make_palette(path: Path, colors: list[tuple[str, str]], title: str) -> None:
    im = Image.new("RGB", (1800, 720), rgb(BG))
    d = ImageDraw.Draw(im)
    d.text((76, 58), title.upper(), font=ft(42, True), fill=rgb(ACCENT))
    x, y, gap = 76, 160, 26
    sw = int((1800 - 152 - gap * (len(colors) - 1)) / len(colors))
    for name, value in colors:
        d.rounded_rectangle((x, y, x + sw, y + 300), radius=26, fill=rgb(value), outline=rgb(LINE), width=4)
        label_color = rgb(TEXT)
        d.text((x, y + 340), name, font=ft(28, True), fill=label_color)
        d.text((x, y + 386), value, font=ft(25), fill=rgb(MUTED))
        x += sw + gap
    im.save(path, optimize=True)


def make_typography(path: Path) -> None:
    im = Image.new("RGB", (1800, 760), rgb(BG))
    d = ImageDraw.Draw(im)
    d.text((76, 55), "TYPOGRAPHY SPECIMEN", font=ft(40, True), fill=rgb(ACCENT))
    d.text((76, 170), "Built for sharp minds.", font=ft(92, display=True), fill=rgb(TEXT))
    d.text((76, 302), "DISPLAY  /  BAHNSCHRIFT PRODUCTION FALLBACK", font=ft(26, True), fill=rgb(ACCENT_2))
    d.text((76, 415), "Clear systems reduce cognitive load and make repeated work easier to trust.", font=ft(38), fill=rgb(TEXT))
    d.text((76, 480), "BODY  /  ARIAL PRODUCTION FALLBACK", font=ft(24, True), fill=rgb(MUTED))
    d.text((76, 610), "INPUT → CONSTRAINT → OUTPUT → CHECK", font=ft(40, True), fill=rgb(ACCENT))
    d.text((1260, 622), "TECHNICAL LABEL", font=ft(22, True), fill=rgb(MUTED))
    im.save(path, optimize=True)


def make_logo_overview(path: Path) -> None:
    src = Image.open(BRAND / "logos" / "primary-logo.png").convert("RGB")
    src.thumbnail((1600, 640), Image.Resampling.LANCZOS)
    im = Image.new("RGB", (1800, 760), rgb(BG))
    d = ImageDraw.Draw(im)
    d.rounded_rectangle((40, 40, 1760, 720), radius=34, fill=rgb(SURFACE), outline=rgb(LINE), width=4)
    im.paste(src, ((1800 - src.width) // 2, (760 - src.height) // 2))
    im.save(path, optimize=True)


def insert_picture_after(paragraph: Paragraph, image: Path, width: float, alt: str) -> None:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    p = Paragraph(new_p, paragraph._parent)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = 0
    p.paragraph_format.space_after = 0
    run = p.add_run()
    shape = run.add_picture(str(image), width=Inches(width))
    shape._inline.docPr.set("descr", alt)
    shape._inline.docPr.set("title", alt)


def paragraph_after_heading(doc: Document, heading: str) -> Paragraph:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == heading:
            if i + 1 >= len(doc.paragraphs):
                raise RuntimeError(f"No paragraph after {heading}")
            return doc.paragraphs[i + 1]
    raise RuntimeError(f"Heading not found: {heading}")


def main() -> None:
    if not DOCX.exists():
        raise SystemExit(f"Missing brand guide: {DOCX}")
    primary = BRAND / "brandbook" / "primary-color-palette.png"
    secondary = BRAND / "brandbook" / "secondary-color-palette.png"
    type_spec = BRAND / "brandbook" / "typography-specimen.png"
    logo_overview = BRAND / "brandbook" / "logo-system-overview.png"
    for path in (primary, secondary, type_spec, logo_overview):
        if path.exists():
            raise SystemExit(f"Refusing to overwrite existing specimen: {path}")
    make_palette(primary, [("Signal Black", BG), ("System Surface", SURFACE), ("Ninja Green", ACCENT), ("Cloud White", TEXT)], "Primary palette")
    make_palette(secondary, [("Insight Cyan", ACCENT_2), ("Steel", MUTED), ("Grid", LINE), ("Deep Surface", SURFACE_2)], "Secondary palette")
    make_typography(type_spec)
    make_logo_overview(logo_overview)

    doc = Document(DOCX)
    rel_targets = [getattr(r, "target_ref", "") for r in doc.part.rels.values()]
    if any("logo-system-overview" in t for t in rel_targets):
        raise SystemExit("Brand guide is already enriched")
    insert_picture_after(paragraph_after_heading(doc, "Logo Concept"), BRAND / "logos" / "icon-only.png", 1.25, "NumberNinjaDesign modular N icon")
    insert_picture_after(paragraph_after_heading(doc, "Logo System"), logo_overview, 5.45, "NumberNinjaDesign primary logo system overview")
    insert_picture_after(paragraph_after_heading(doc, "Primary Color Palette"), primary, 5.55, "Primary palette with names and hexadecimal values")
    insert_picture_after(paragraph_after_heading(doc, "Secondary and Accent Palette"), secondary, 5.55, "Secondary palette with names and hexadecimal values")
    insert_picture_after(paragraph_after_heading(doc, "Typography System"), type_spec, 5.55, "Production typography specimen and fallback hierarchy")
    insert_picture_after(paragraph_after_heading(doc, "Etsy Branding"), BRAND / "banners" / "etsy-banner.png", 5.55, "Etsy big banner for NumberNinjaDesign")
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    main()
